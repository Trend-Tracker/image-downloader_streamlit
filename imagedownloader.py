"""
Images to ZIP + PPT maker - Streamlit app (low-memory version).

Three ways to use it:
  1. Upload Excel  -> downloads every image column, files named  <StoreCode>_<ColumnName>[_n].jpg
  2. Paste links   -> one URL per line, get a ZIP
  3. Make PPT      -> upload Excel, get a PowerPoint: one slide per outlet
                      (Store front | Before deployment | After deployment)
                      with a blue header: client store code + store name

Why this version does not crash on ~3,000 images:
  * Images are written straight into ZIP files on disk and thrown away from memory.
  * Output is split into parts (e.g. 200 images per ZIP, 40 outlets per PPT).
  * Only one part is held in memory at a time (also when you click Download).
  * Photos are decoded at reduced size and shrunk before they go into the PPT.

Run locally:
    pip install streamlit openpyxl python-pptx pillow
    streamlit run imagedownloader.py
"""
import base64
import gc
import io
import math
import mimetypes
import re
import shutil
import tempfile
import zipfile
from datetime import date, datetime
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from urllib.parse import parse_qs, unquote, urlparse
import urllib.request

import streamlit as st
from openpyxl import load_workbook
from PIL import Image, ImageOps
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

HEADERS = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) Chrome/124.0 Safari/537.36"}
URL_RE = re.compile(r"https?://[^\s,\"'<>]+", re.I)

MAX_DOWNLOAD_BYTES = 40 * 1024 * 1024   # skip a single file bigger than this
FETCH_CHUNK = 24                        # how many images are in flight at once (keeps RAM low)
DEFAULT_IMAGES_PER_ZIP = 200
DEFAULT_OUTLETS_PER_PPT = 40

# Built-in logo (base64 JPG/PNG). Leave empty to show the text "Channelplay" on slides,
# or upload your logo inside the app.
DEFAULT_LOGO_B64 = ""


# --------------------------------------------------------------------------
# Temp-folder helpers (big files live on disk, not in RAM)
# --------------------------------------------------------------------------
def fresh_workdir(key: str) -> Path:
    """New temp folder for this kind of output; the previous one is deleted."""
    old = st.session_state.pop(f"{key}_dir", None)
    if old:
        shutil.rmtree(old, ignore_errors=True)
    path = tempfile.mkdtemp(prefix=f"{key}_")
    st.session_state[f"{key}_dir"] = path
    return Path(path)


def show_download(parts, key, label, mime):
    """One download button. With several parts, pick the part first (only one part is loaded in RAM)."""
    if not parts:
        return
    if len(parts) > 1:
        st.info(f"The output is split into {len(parts)} parts to keep the app fast and stable. "
                f"Pick a part, download it, then pick the next one.")
        i = st.selectbox(
            "Part to download", list(range(len(parts))), key=f"{key}_part",
            format_func=lambda n: (f"Part {n + 1} of {len(parts)} - {parts[n]['count']} {parts[n]['unit']}"
                                   f" ({parts[n]['mb']:.0f} MB)"),
        )
    else:
        i = 0
    part = parts[i]
    path = Path(part["path"])
    if not path.exists():
        st.error("This file is no longer on the server (the app restarted). Please create it again.")
        return
    with open(path, "rb") as f:
        st.download_button(label if len(parts) == 1 else f"{label} - part {i + 1}", data=f.read(),
                           file_name=part["name"], mime=mime, type="primary", key=f"{key}_dl_{i}")


# --------------------------------------------------------------------------
# Excel / download helpers
# --------------------------------------------------------------------------
def parse_links(text: str) -> list[str]:
    seen, links = set(), []
    for token in re.split(r"[\s,]+", text):
        if re.match(r"^https?://\S+$", token, re.I) and token not in seen:
            seen.add(token)
            links.append(token)
    return links


def clean(text: str) -> str:
    """Make a string safe for use in a file name."""
    return re.sub(r"[^\w\-]+", "_", str(text).strip()).strip("_")


def expand_gallery(url: str) -> list[str]:
    """Gallery links (.../gallery?i=<base64>) hold the real image URLs, one per line."""
    try:
        i = parse_qs(urlparse(url).query)["i"][0]
        i += "=" * (-len(i) % 4)
        decoded = base64.urlsafe_b64decode(i).decode()
        return URL_RE.findall(decoded) or [url]
    except Exception:
        return [url]


def cell_urls(cell) -> list[str]:
    """All direct image URLs in one Excel cell (hyperlink target, plain text, or gallery link)."""
    candidates = []
    if cell.hyperlink is not None and cell.hyperlink.target:
        candidates.append(cell.hyperlink.target)
    val = cell.value
    if isinstance(val, str):
        candidates += URL_RE.findall(val)  # plain URL text or =HYPERLINK("url","text")

    urls, seen = [], set()
    for c in candidates:
        for u in (expand_gallery(c) if "/gallery" in c and "i=" in c else [c]):
            if u not in seen:
                seen.add(u)
                urls.append(u)
    return urls


@st.cache_resource(show_spinner=False, max_entries=2)
def get_sheetnames(file_bytes: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(file_bytes), read_only=True)
    names = list(wb.sheetnames)
    wb.close()
    return names


@st.cache_resource(show_spinner="Reading Excel...", max_entries=2)
def read_workbook(file_bytes: bytes, sheet_name=None):
    """Cached, so the workbook is parsed once and not on every click."""
    wb = load_workbook(io.BytesIO(file_bytes))  # not read_only: hyperlinks need full mode
    ws = wb[sheet_name] if sheet_name else wb.active
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    return wb, ws, headers


def build_jobs(ws, headers, code_col, image_cols, rows=None):
    """One job per image: (url, file stem, store code, column name)."""
    code_idx = headers.index(code_col)
    img_idx = {name: headers.index(name) for name in image_cols}

    jobs, used = [], {}  # used: file stem -> url
    for row in range(2, ws.max_row + 1):
        if rows is not None and row not in rows:
            continue
        code = ws.cell(row=row, column=code_idx + 1).value
        code = clean(code) if code not in (None, "") else f"row{row}"
        for col_name, idx in img_idx.items():
            urls = cell_urls(ws.cell(row=row, column=idx + 1))
            for n, url in enumerate(urls, start=1):
                stem = f"{code}_{clean(col_name)}"
                if len(urls) > 1:
                    stem += f"_{n}"
                if stem in used:
                    if used[stem] == url:
                        continue                      # exact same photo repeated in another row
                    stem += f"_r{row}"                # same name, different photo -> keep both
                used[stem] = url
                jobs.append({"url": url, "stem": stem, "code": code, "column": col_name})
    return jobs


def ext_for(url: str, content_type: str) -> str:
    suffix = Path(urlparse(url).path).suffix
    if re.fullmatch(r"\.[A-Za-z0-9]{2,5}", suffix or ""):
        return suffix.lower()
    return mimetypes.guess_extension((content_type or "").split(";")[0]) or ".jpg"


def fetch_raw(url: str):
    """Download one URL -> (bytes, content type). Retries once on network errors."""
    last = None
    for _ in range(2):
        try:
            req = urllib.request.Request(url, headers=HEADERS)
            with urllib.request.urlopen(req, timeout=30) as resp:
                data = resp.read(MAX_DOWNLOAD_BYTES + 1)
                ctype = resp.headers.get("Content-Type", "")
            if len(data) > MAX_DOWNLOAD_BYTES:
                raise ValueError("file is too large")
            if ctype.startswith(("text/", "application/json")):
                raise ValueError(f"got a web page instead of an image ({ctype})")
            return data, ctype
        except ValueError:
            raise                                     # retrying will not help
        except Exception as err:
            last = err
    raise last


def fetch(job):
    try:
        data, ctype = fetch_raw(job["url"])
        return {**job, "name": job["stem"] + ext_for(job["url"], ctype), "data": data, "error": None}
    except Exception as err:
        return {**job, "name": None, "data": None, "error": str(err)}


def run_jobs(jobs, zip_name, per_zip):
    """Download images and write them straight into ZIP files on disk, `per_zip` images per file."""
    workdir = fresh_workdir("zip")
    base = re.sub(r"[^\w\- ]+", "_", zip_name.strip() or "images")
    n_parts = max(1, math.ceil(len(jobs) / per_zip))
    progress = st.progress(0.0, text=f"Fetching 0 of {len(jobs)}...")

    report, parts, done = [], [], 0
    with ThreadPoolExecutor(max_workers=4) as pool:
        for p in range(n_parts):
            part_jobs = jobs[p * per_zip:(p + 1) * per_zip]
            fname = f"{base}.zip" if n_parts == 1 else f"{base}_part{p + 1}.zip"
            path = workdir / fname
            added = 0
            with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
                for i in range(0, len(part_jobs), FETCH_CHUNK):
                    for r in pool.map(fetch, part_jobs[i:i + FETCH_CHUNK]):
                        if not r["error"]:
                            zf.writestr(r["name"], r["data"])
                            added += 1
                        report.append({
                            "Store Code": r.get("code", ""),
                            "Column": r.get("column", ""),
                            "Link": r["url"],
                            "Status": "Added" if not r["error"] else "Failed",
                            "Details": r["name"] or r["error"],
                        })
                        r["data"] = None              # free the image right away
                        done += 1
                        progress.progress(done / len(jobs),
                                          text=f"Fetching {done} of {len(jobs)}  (ZIP part {p + 1} of {n_parts})...")
            if added:
                parts.append({"path": str(path), "name": fname, "count": added, "unit": "images",
                              "mb": path.stat().st_size / 1024 / 1024})
            else:
                path.unlink(missing_ok=True)
            gc.collect()
    progress.empty()

    st.session_state["zip_result"] = {
        "parts": parts, "report": report,
        "good": sum(p["count"] for p in parts), "total": len(jobs),
    }


# --------------------------------------------------------------------------
# Date filter
# --------------------------------------------------------------------------
def parse_date(v):
    """Excel date cell -> datetime.date, or None. Handles real dates, 2026-10-17 and 17/10/2026 (day first)."""
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = str(v or "").strip()
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if m:
        y, mo, d = map(int, m.groups())
    else:
        m = re.match(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})", t)
        if not m:
            return None
        d, mo, y = map(int, m.groups())
    try:
        return date(y, mo, d)
    except ValueError:
        return None


def date_filter_ui(ws, headers, key):
    """Shows the date filter. Returns the set of Excel row numbers to process, or None for 'all rows'."""
    named = [h for h in headers if h]
    NO = "(no date filter)"
    default = pick(named, "check-in date", default=-1)
    col = st.selectbox("Date filter column", [NO] + named, index=default + 1 if default >= 0 else 0,
                       key=f"{key}_datecol")
    if col == NO:
        return None
    idx = headers.index(col) + 1
    dates = {r: parse_date(ws.cell(row=r, column=idx).value) for r in range(2, ws.max_row + 1)}
    valid = [d for d in dates.values() if d]
    if not valid:
        st.warning(f"No readable dates found in “{col}”, so the date filter is off.")
        return None
    lo, hi = min(valid), max(valid)
    c1, c2 = st.columns(2)
    start = c1.date_input("From date", value=lo, min_value=lo, max_value=hi, key=f"{key}_from_{lo}_{hi}")
    end = c2.date_input("To date", value=hi, min_value=lo, max_value=hi, key=f"{key}_to_{lo}_{hi}")
    allowed = {r for r, d in dates.items() if d and start <= d <= end}
    unreadable = sum(1 for d in dates.values() if d is None)
    st.caption(f"{len(allowed)} of {len(dates)} rows are inside {start:%d %b} – {end:%d %b %Y}"
               + (f"; {unreadable} rows have no readable date and are skipped" if unreadable else ""))
    return allowed


# --------------------------------------------------------------------------
# PPT builder
# --------------------------------------------------------------------------
BLUE = RGBColor(0x00, 0x70, 0xC0)
GREY = RGBColor(0x7F, 0x7F, 0x7F)
LIGHT = RGBColor(0xF2, 0xF2, 0xF2)
SLIDE_W, SLIDE_H = 13.333, 7.5       # inches (16:9)
BAR_W, BAR_H = 10.7, 0.65            # blue header bar (logo sits to its right, on white)
MARGIN = 0.35
COL_GAP = 0.25
COLUMNS = [("front", "Store Front"), ("before", "Before Deployment"), ("after", "After Deployment")]


def prep_image(data: bytes, max_side=1200):
    """Fix phone-photo rotation, shrink and re-encode so the PPT stays small. -> (jpeg bytes, (w, h))"""
    try:
        im = Image.open(io.BytesIO(data))
        im.draft("RGB", (max_side, max_side))          # JPEG: decode at reduced size = far less RAM
        im = ImageOps.exif_transpose(im).convert("RGB")
        im.thumbnail((max_side, max_side))
        buf = io.BytesIO()
        im.save(buf, "JPEG", quality=80, optimize=True)
        return buf.getvalue(), im.size
    except Exception:
        return None


def boxes_for(n, x, y, w, h, gap=0.12):
    """Where each of n images goes inside a column box."""
    if n <= 1:
        return [(x, y, w, h)]
    if n == 2:
        bh = (h - gap) / 2
        return [(x, y, w, bh), (x, y + bh + gap, w, bh)]
    if n == 3:
        bh, bw = (h - gap) / 2, (w - gap) / 2
        return [(x, y, w, bh), (x, y + bh + gap, bw, bh), (x + bw + gap, y + bh + gap, bw, bh)]
    rows = math.ceil(n / 2)
    bh, bw = (h - gap * (rows - 1)) / rows, (w - gap) / 2
    return [(x + (i % 2) * (bw + gap), y + (i // 2) * (bh + gap), bw, bh) for i in range(n)]


def _rect(slide, x, y, w, h, fill):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.fill.background()
    shp.shadow.inherit = False
    for ref in shp._element.xpath(".//a:effectRef"):   # also drop the theme's shadow reference
        ref.set("idx", "0")
    return shp


def _text(shape_or_tf, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    tf = shape_or_tf.text_frame if hasattr(shape_or_tf, "text_frame") else shape_or_tf
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.name = "Calibri"
    r.font.color.rgb = color
    return tf


def add_logo(slide, logo_bytes):
    h = 0.5
    try:
        w_px, h_px = Image.open(io.BytesIO(logo_bytes)).size
        w = h * w_px / h_px
        slide.shapes.add_picture(io.BytesIO(logo_bytes), Inches(SLIDE_W - MARGIN + 0.1 - w), Inches(0.08),
                                 Inches(w), Inches(h))
    except Exception:
        tb = slide.shapes.add_textbox(Inches(BAR_W + 0.2), Inches(0.08), Inches(SLIDE_W - BAR_W - 0.4), Inches(h))
        tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(tb, "Channelplay", 22, BLUE, bold=True, align=PP_ALIGN.RIGHT)


def add_header(slide, title, logo_bytes):
    bar = _rect(slide, 0, 0, BAR_W, BAR_H, BLUE)
    tf = bar.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.25)
    size = 20 if len(title) <= 58 else 18 if len(title) <= 70 else 15
    _text(tf, title, size, RGBColor(255, 255, 255))
    add_logo(slide, logo_bytes)


def add_welcome(prs, title, subtitle, logo_bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(slide, logo_bytes)
    band = _rect(slide, 0, 2.7, SLIDE_W, 1.5, BLUE)
    band.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    band.text_frame.word_wrap = True
    _text(band, title, 36, RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)
    if subtitle:
        tb = slide.shapes.add_textbox(Inches(1), Inches(4.4), Inches(SLIDE_W - 2), Inches(0.7))
        tb.text_frame.word_wrap = True
        _text(tb, subtitle, 20, GREY, align=PP_ALIGN.CENTER)


def add_thanks(prs, logo_bytes):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_logo(slide, logo_bytes)
    band = _rect(slide, 0, 2.7, SLIDE_W, 1.5, BLUE)
    band.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(band, "Thank You", 40, RGBColor(255, 255, 255), bold=True, align=PP_ALIGN.CENTER)


def add_image_box(slide, img, box):
    """Place one prepared image (or a 'no image' placeholder) fitted inside box."""
    bx, by, bw, bh = box
    if img is None:
        ph = _rect(slide, bx, by, bw, bh, LIGHT)
        ph.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(ph, "Image not available", 12, GREY, align=PP_ALIGN.CENTER)
        return
    data, (iw, ih) = img
    scale = min(bw / iw, bh / ih)
    w, h = iw * scale, ih * scale
    pic = slide.shapes.add_picture(io.BytesIO(data), Inches(bx + (bw - w) / 2), Inches(by + (bh - h) / 2),
                                   Inches(w), Inches(h))
    pic.line.color.rgb = GREY
    pic.line.width = Pt(0.75)


def add_outlet_slide(prs, entry, logo_bytes, images, active, show_titles=False):
    """entry: {title, front, before, after} (url lists); images: url -> prepared image or None.
    active: the (key, label) columns that are switched on; they share the slide width equally."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header(slide, entry["title"], logo_bytes)

    n = len(active)
    col_w = (SLIDE_W - 2 * MARGIN - (n - 1) * COL_GAP) / n
    if show_titles:
        label_y, label_h = BAR_H + 0.2, 0.35
        area_y = label_y + label_h + 0.12
    else:
        area_y = BAR_H + 0.3
    area_h = SLIDE_H - area_y - 0.3
    for i, (key, label) in enumerate(active):
        x = MARGIN + i * (col_w + COL_GAP)
        if show_titles:
            tb = slide.shapes.add_textbox(Inches(x), Inches(label_y), Inches(col_w), Inches(label_h))
            tb.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            _text(tb, label, 14, BLUE, bold=True, align=PP_ALIGN.CENTER)
            _rect(slide, x, label_y + label_h, col_w, 0.03, BLUE)

        urls = entry[key]
        if not urls:
            add_image_box(slide, None, (x, area_y, col_w, area_h))
            continue
        for url, box in zip(urls, boxes_for(len(urls), x, area_y, col_w, area_h)):
            add_image_box(slide, images.get(url), box)


def build_deck(entries, images, logo_bytes, welcome_sub="", active=COLUMNS, show_titles=False):
    """Welcome slide per campaign (name from the Excel) + its outlet slides, then a closing Thank You slide."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(SLIDE_W), Inches(SLIDE_H)
    groups = {}
    for e in entries:
        groups.setdefault(e["campaign"], []).append(e)
    for campaign, items in groups.items():
        if campaign:
            add_welcome(prs, campaign, welcome_sub, logo_bytes)
        for e in items:
            add_outlet_slide(prs, e, logo_bytes, images, active, show_titles)
    add_thanks(prs, logo_bytes)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_entries(ws, headers, cols, rows=None, active_keys=None):
    """One entry per Excel row (identical rows skipped).

    cols maps: front/before/after (image columns), code, client, name, city, state, campaign.
    Header text = "<client code> – <store name> – <city>, <state>"; campaign = welcome-slide title.
    """
    idx = {h: i + 1 for i, h in enumerate(headers) if h}
    entries, seen = [], set()
    for row in range(2, ws.max_row + 1):
        if rows is not None and row not in rows:
            continue

        def text(key):
            h = cols.get(key)
            if not h or h not in idx:
                return ""
            v = ws.cell(row=row, column=idx[h]).value
            return "" if v is None else str(v).strip()

        urls = {k: (cell_urls(ws.cell(row=row, column=idx[cols[k]])) if active_keys is None or k in active_keys else [])
                for k, _ in COLUMNS}
        place = ", ".join(p for p in (text("city"), text("state")) if p)
        title = " – ".join(p for p in (text("client"), text("name"), place) if p) or f"Row {row}"
        code = text("code") or f"row{row}"
        campaign = text("campaign")
        key = (code, title, campaign, tuple(urls["front"]), tuple(urls["before"]), tuple(urls["after"]))
        if key in seen:
            continue
        seen.add(key)
        entries.append({"code": code, "name": text("name"), "title": title, "campaign": campaign, **urls})
    return entries


def fetch_prepared(url):
    """Download + shrink one image inside the worker thread, so only the small version is kept."""
    try:
        data, _ = fetch_raw(url)
        prepared = prep_image(data)
        if prepared is None:
            return url, None, "not a readable image"
        return url, prepared, None
    except Exception as err:
        return url, None, str(err)


def download_all(urls, label="Images", workers=6):
    """url -> prepared image (or None), for one part only. Returns (images, failures)."""
    progress = st.progress(0.0, text=f"{label}: downloading 0 of {len(urls)} images...")
    images, failures, done = {}, [], 0
    with ThreadPoolExecutor(max_workers=workers) as pool:
        for i in range(0, len(urls), FETCH_CHUNK):
            for url, prepared, err in pool.map(fetch_prepared, urls[i:i + FETCH_CHUNK]):
                images[url] = prepared
                if prepared is None:
                    failures.append({"Link": url, "Reason": err or "not a readable image"})
                done += 1
                progress.progress(done / max(len(urls), 1),
                                  text=f"{label}: downloading {done} of {len(urls)} images...")
    progress.empty()
    return images, failures


def create_ppt_parts(entries, active, logo_bytes, w_sub, show_titles, per_outlet_zip, outlets_per_part):
    """Build the PPT(s) part by part. Images of one part are downloaded, used, then dropped from memory."""
    workdir = fresh_workdir("ppt")

    # split the outlets into parts (in per-outlet mode an outlet code is never split across parts)
    chunks = []
    if per_outlet_zip:
        by_code = {}
        for e in entries:
            by_code.setdefault(e["code"], []).append(e)
        cur = []
        for items in by_code.values():
            cur += items
            if len(cur) >= outlets_per_part:
                chunks.append(cur)
                cur = []
        if cur:
            chunks.append(cur)
    else:
        chunks = [entries[i:i + outlets_per_part] for i in range(0, len(entries), outlets_per_part)]

    base = "outlet_ppts" if per_outlet_zip else "outlet_report"
    ext = ".zip" if per_outlet_zip else ".pptx"
    parts, failures, total_urls = [], [], 0
    note = st.empty()
    for n, chunk in enumerate(chunks, start=1):
        label = f"Part {n} of {len(chunks)}"
        urls = list(dict.fromkeys(u for e in chunk for k, _ in active for u in e[k]))
        images, fails = download_all(urls, label)
        failures += fails
        total_urls += len(urls)

        note.info(f"{label}: building PowerPoint...")
        fname = f"{base}{ext}" if len(chunks) == 1 else f"{base}_part{n}{ext}"
        path = workdir / fname
        if per_outlet_zip:
            groups = {}
            for e in chunk:
                groups.setdefault(e["code"], []).append(e)
            with zipfile.ZipFile(path, "w", zipfile.ZIP_STORED) as zf:
                for code, items in groups.items():
                    zf.writestr(f"{clean(code)}_{clean(items[0]['name'])}.pptx".replace("_.pptx", ".pptx"),
                                build_deck(items, images, logo_bytes, w_sub, active, show_titles))
        else:
            path.write_bytes(build_deck(chunk, images, logo_bytes, w_sub, active, show_titles))
        parts.append({"path": str(path), "name": fname, "count": len(chunk), "unit": "outlets",
                      "mb": path.stat().st_size / 1024 / 1024})
        del images
        gc.collect()
    note.empty()
    return parts, failures, total_urls


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
def pick(headers, *needles, default=0):
    for i, h in enumerate(headers):
        if h and all(n in h.lower() for n in needles):
            return i
    return default


def default_logo() -> bytes:
    """Built-in logo. If it cannot be decoded, return b"" and the slides show the text "Channelplay" instead."""
    try:
        raw = base64.b64decode(DEFAULT_LOGO_B64 + "=" * (-len(DEFAULT_LOGO_B64) % 4))
        Image.open(io.BytesIO(raw)).load()
        return raw
    except Exception:
        return b""


def per_zip_input(key):
    return st.number_input(
        "Images per ZIP part", min_value=20, max_value=1000, value=DEFAULT_IMAGES_PER_ZIP, step=20, key=key,
        help="Large jobs are split into several ZIP files so the app does not run out of memory. "
             "If the app still crashes, lower this number.")


def main():
    st.set_page_config(page_title="Images to ZIP & PPT Maker", page_icon="🗜️", layout="centered")

    tab_excel, tab_links, tab_ppt = st.tabs(["📄 Upload Excel", "🔗 Paste links", "📊 Make PPT"])

    # ---------------- Excel -> ZIP ----------------
    with tab_excel:
        st.title("Images to ZIP")
        st.write(
            "Upload your Excel. Every image column is downloaded, including the "
            "“N photos — open” gallery cells. Files are named **StoreCode_ColumnName.jpg**."
        )
        up = st.file_uploader("Excel file", type=["xlsx", "xlsm"], key="up_zip")
        if up:
            file_bytes = up.getvalue()
            sheets = get_sheetnames(file_bytes)
            sheet = st.selectbox("Sheet", sheets, key="sh_zip") if len(sheets) > 1 else sheets[0]
            wb, ws, headers = read_workbook(file_bytes, sheet)
            named = [h for h in headers if h]

            code_col = st.selectbox("Outlet code column", named, index=pick(named, "store code"), key="code_zip")
            auto_img = [h for h in named if "image" in h.lower()]
            image_cols = st.multiselect("Image columns", named, default=auto_img)

            rows = date_filter_ui(ws, headers, "zip")
            jobs = build_jobs(ws, headers, code_col, image_cols, rows) if image_cols else []
            st.caption(f"{len(jobs)} image{'s' if len(jobs) != 1 else ''} found in {ws.max_row - 1 if rows is None else len(rows)} rows")
            if jobs:
                with st.expander("Preview file names"):
                    st.dataframe([{"File": j["stem"] + ".jpg", "Link": j["url"]} for j in jobs[:500]], hide_index=True)
                    if len(jobs) > 500:
                        st.caption(f"Showing the first 500 of {len(jobs)}.")

            c1, c2 = st.columns(2)
            zip_name_x = c1.text_input("ZIP file name", value="images", key="zip_x")
            with c2:
                per_zip_x = per_zip_input("per_zip_x")
            if st.button("Prepare ZIP", type="primary", disabled=not jobs, key="go_x"):
                st.session_state.pop("zip_result", None)
                run_jobs(jobs, zip_name_x, int(per_zip_x))

    # ---------------- Paste links -> ZIP ----------------
    with tab_links:
        st.title("Images to ZIP")
        st.write("Paste image links, one per line. Get them all in a single ZIP file.")
        raw = st.text_area("Image URLs", height=220,
                           placeholder="https://example.com/photo1.jpg\nhttps://example.com/photo2.png")
        links = parse_links(raw)

        col1, col2, col3 = st.columns([3, 2, 1], vertical_alignment="bottom")
        zip_name_l = col1.text_input("ZIP file name", value="images", key="zip_l")
        with col2:
            per_zip_l = per_zip_input("per_zip_l")
        col3.caption(f"{len(links)} link{'s' if len(links) != 1 else ''}")

        if st.button("Prepare ZIP", type="primary", disabled=not links, key="go_l"):
            st.session_state.pop("zip_result", None)
            jobs = []
            for i, u in enumerate(links):
                base = re.sub(r"[^\w.\-]+", "_", unquote(Path(urlparse(u).path).name) or "image")
                jobs.append({"url": u, "stem": f"{i + 1:03d}_{Path(base).stem}", "code": "", "column": ""})
            run_jobs(jobs, zip_name_l, int(per_zip_l))

    # ---------------- ZIP result (shared by the two tabs above) ----------------
    result = st.session_state.get("zip_result")
    if result:
        st.divider()
        if result["parts"]:
            msg = f"{result['good']} of {result['total']} images are ready."
            (st.success if result["good"] == result["total"] else st.warning)(msg)
            show_download(result["parts"], "zip", "Download ZIP", "application/zip")
        else:
            st.error("No images could be downloaded. Check that the links point directly to image files.")
        st.dataframe(result["report"], width="stretch", hide_index=True)

    # ---------------- Excel -> PPT ----------------
    with tab_ppt:
        st.title("✨ Outlet Deck Maker")
        st.subheader("From Excel to a client-ready PPT in one click")
        st.write("Upload your Excel to get a PowerPoint with **one slide per outlet**: "
                 "store front, then before deployment, then after deployment.")
        up = st.file_uploader("Excel file", type=["xlsx", "xlsm"], key="up_ppt")
        if up:
            file_bytes = up.getvalue()
            sheets = get_sheetnames(file_bytes)
            sheet = st.selectbox("Sheet", sheets, key="sh_ppt") if len(sheets) > 1 else sheets[0]
            wb, ws, headers = read_workbook(file_bytes, sheet)
            named = [h for h in headers if h]

            st.markdown("**Images to include** (switch off what you don't want on the slides)")
            c1, c2, c3 = st.columns(3)
            front_col = c1.selectbox("Store front column", named, index=pick(named, "front"))
            on_front = c1.toggle("Store front", value=True)
            before_col = c2.selectbox("Before column", named, index=pick(named, "before"))
            on_before = c2.toggle("Before deployment", value=True)
            after_col = c3.selectbox("After column", named, index=pick(named, "after"))
            on_after = c3.toggle("After deployment", value=True)
            active = [(k, l) for (k, l), on in zip(COLUMNS, (on_front, on_before, on_after)) if on]
            if not active:
                st.warning("Switch on at least one image type.")
            show_titles = st.checkbox("Show column titles on slides (Store Front / Before / After)", value=False)

            rows = date_filter_ui(ws, headers, "ppt")

            code_col = st.selectbox("Outlet code column (used for file names)", named,
                                    index=pick(named, "store code"), key="code_ppt")

            st.markdown("**Blue header text:** client code – store name – city, state")
            h1, h2, h3, h4 = st.columns(4)
            NONE = "(none)"
            opts = [NONE] + named
            def opt_idx(*needles):
                """Index into opts for the first header containing all needles, else 0 = (none)."""
                i = pick(named, *needles, default=-1)
                return i + 1 if i >= 0 else 0

            client_col = h1.selectbox("Client code", opts, index=opt_idx("client", "code"))
            name_col = h2.selectbox("Store name", opts, index=opt_idx("store name"))
            city_col = h3.selectbox("City", opts, index=opt_idx("city"))
            state_col = h4.selectbox("State", opts, index=opt_idx("state"))

            st.markdown("**Welcome slide**")
            camp_col = st.selectbox("Campaign column (its value becomes the welcome-slide title)", opts,
                                    index=opt_idx("campaign", "name"))
            w_sub = st.text_input("Subtitle (optional)", value="")
            logo_up = st.file_uploader("Logo (optional – a sharper Channelplay logo, PNG/JPG)",
                                       type=["png", "jpg", "jpeg"])
            logo_bytes = logo_up.getvalue() if logo_up else default_logo()

            o1, o2 = st.columns([3, 2])
            mode = o1.radio("Output", ["One PPT with all outlets", "One PPT per outlet (ZIP)"], horizontal=True)
            per_part = o2.number_input(
                "Outlets per file", min_value=10, max_value=200, value=DEFAULT_OUTLETS_PER_PPT, step=10,
                help="Big jobs are split into several files so the app does not run out of memory. "
                     "If the app still crashes, lower this number.")

            cols = {"front": front_col, "before": before_col, "after": after_col, "code": code_col,
                    "client": None if client_col == NONE else client_col,
                    "name": None if name_col == NONE else name_col,
                    "city": None if city_col == NONE else city_col,
                    "state": None if state_col == NONE else state_col,
                    "campaign": None if camp_col == NONE else camp_col}
            entries = build_entries(ws, headers, cols, rows, {k for k, _ in active}) if active else []
            if entries:
                st.caption("Header preview: " + entries[0]["title"]
                           + (f"   |   Welcome slide: {entries[0]['campaign']}" if entries[0]["campaign"] else ""))
            all_urls = list(dict.fromkeys(u for e in entries for k, _ in active for u in e[k]))
            n_files = math.ceil(len(entries) / int(per_part)) if entries else 0
            st.caption(f"{len(entries)} outlet slide{'s' if len(entries) != 1 else ''}, {len(all_urls)} images"
                       + (f"  →  {n_files} file{'s' if n_files != 1 else ''}" if n_files > 1 else ""))

            if st.button("Create PPT", type="primary", disabled=not entries, key="go_ppt"):
                st.session_state.pop("ppt_result", None)
                per_outlet_zip = not mode.startswith("One PPT with")
                parts, failures, total = create_ppt_parts(entries, active, logo_bytes, w_sub, show_titles,
                                                          per_outlet_zip, int(per_part))
                st.session_state["ppt_result"] = {
                    "parts": parts, "failures": failures, "total": total,
                    "mime": "application/zip" if per_outlet_zip else
                            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    "label": "Download ZIP of PPTs" if per_outlet_zip else "Download PPT",
                }

        pres = st.session_state.get("ppt_result")
        if pres:
            st.divider()
            ok = pres["total"] - len(pres["failures"])
            (st.success if not pres["failures"] else st.warning)(
                f"PPT ready. {ok} of {pres['total']} images placed."
                + ("" if not pres["failures"] else " Missing ones show “Image not available”."))
            show_download(pres["parts"], "ppt", pres["label"], pres["mime"])
            if pres["failures"]:
                st.dataframe(pres["failures"], width="stretch", hide_index=True)


if __name__ == "__main__":
    main()
